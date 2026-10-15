#!/usr/bin/env python3
"""Builds the Nomad Compass Grant Portfolio Excel upload template.

Output: ~/workspace/your_files/nomad-compass/nomad-compass-grant-upload-template.xlsx

Sheets:
  1. Instructions - fill, save as CSV UTF-8, upload via Grant Tracking -> Bulk Grant Upload.
  2. Grants - fillable table matching the bulk importer in
     components/GrantCsvImport.tsx / src/lib/grantCsvImport.ts.

Column enums mirror the app's validation:
  Status: active | pending | completed        (default: active)
  Restriction: restricted | unrestricted
  Report Frequency: monthly | quarterly | semiannual | annual | none
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = os.path.expanduser("~/workspace/your_files/nomad-compass")
OUT_FILE = os.path.join(OUT_DIR, "nomad-compass-grant-upload-template.xlsx")

TEAL = "0F766E"
TEAL_DARK = "115E59"
BRASS = "B45309"
INK = "1F2937"
GREY = "6B7280"
WHITE = "FFFFFF"
ROW_ALT = "F9FAFB"
EXAMPLE_BG = "F3F4F6"

thin = Side(style="thin", color="D1D5DB")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADERS = [
    "Grant Name",
    "Funder",
    "Award Amount",
    "Amount Spent",
    "Start Date",
    "End Date",
    "Status",
    "Restriction",
    "Report Frequency",
    "Program",
]
COL_WIDTHS = [30, 26, 16, 16, 14, 14, 13, 14, 18, 24]

EXAMPLES = [
    ["Youth Workforce Expansion", "Ford Foundation", 150000, 42500, "2026-01-01", "2026-12-31", "active", "restricted", "quarterly", "Youth Workforce"],
    ["Water Access Initiative", "Bloomberg Philanthropies", 75000, 0, "2026-03-01", "2027-02-28", "pending", "unrestricted", "annual", ""],
]

INSTRUCTIONS_BODY = [
    ("What this file is for",
     "This template feeds the Grant Tracking section of Nomad Compass. Fill it in — one row per grant — "
     "save it as a CSV file, and upload it under Grant Tracking → Bulk Grant Upload. Each row becomes a new "
     "grant in your portfolio. KPIs and subgrantees are added afterwards, in each grant's detail view."),
    ("Step 1 — Fill in the Grants sheet",
     "• Replace or delete the two grey EXAMPLE rows — they are placeholders only.\n"
     "• Grant Name, Funder, and Award Amount are required. Everything else is optional.\n"
     "• Amounts are plain numbers — no $ or commas needed (they are accepted anyway).\n"
     "• Dates use YYYY-MM-DD (e.g. 2026-01-01). MM/DD/YYYY also works.\n"
     "• Status: active, pending, or completed. Leave blank for active.\n"
     "• Restriction: restricted or unrestricted. Leave blank if unknown.\n"
     "• Report Frequency: monthly, quarterly, semiannual, annual, or none.\n"
     "• Program: must match an existing program name exactly (check Programs first). "
     "If it doesn't match, the grant still imports — it just stays unassigned."),
    ("Step 2 — Save as CSV",
     "In Excel: File → Save As → choose “CSV UTF-8 (Comma delimited) (*.csv)”.\n"
     "In Google Sheets: File → Download → Comma Separated Values (.csv).\n"
     "The file must end in .csv — that is the only format the uploader accepts."),
    ("Step 3 — Upload to Nomad Compass",
     "1. Sign in and open Grant Tracking from the sidebar.\n"
     "2. In the “Bulk Grant Upload” panel, drag your .csv file in (or click to browse).\n"
     "3. Check the column mapping — Nomad Compass guesses it automatically; fix any dropdown that "
     "looks wrong. Grant Name, Funder and Award Amount must be mapped.\n"
     "4. Review the preview: rows with errors are listed and will be skipped; possible duplicates "
     "are flagged but still import.\n"
     "5. Click Import N Grants and watch the progress. Add KPIs and subgrantees afterwards in each "
     "grant's detail view."),
    ("Tips for clean uploads",
     "• Keep Funder names consistent (“Ford Foundation” vs “Ford Fdn” will look like two funders).\n"
     "• One header row at the top, no merged cells, no extra title rows above the headers.\n"
     "• You can rename the column headers if you like — Nomad Compass matches them by meaning, "
     "but the template headers work as-is.\n"
     "• Only Admins and Grant Coordinators can import grants; other roles won't see the upload panel."),
]

wb = Workbook()

# ---------------------------------------------------------------- Instructions
ws = wb.active
ws.title = "Instructions"
ws.sheet_properties.pageSetUpPr.fitToPage = True

ws["A1"].value = "Nomad Compass — Grant Upload Template"
ws["A1"].font = Font(size=18, bold=True, color=WHITE)
ws["A1"].fill = PatternFill("solid", fgColor=TEAL)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 42
ws.column_dimensions["A"].width = 110
ws.sheet_view.showGridLines = False

row = 3
for heading, body in INSTRUCTIONS_BODY:
    h = ws[f"A{row}"]
    h.value = heading
    h.font = Font(size=13, bold=True, color=TEAL_DARK)
    h.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26
    row += 1
    b = ws[f"A{row}"]
    b.value = body
    b.font = Font(size=11, color=INK)
    b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = max(70, 18 * (len(body) // 100 + body.count("\n") + 1))
    row += 2

ws[f"A{row}"].value = ("Questions? Only Admins and Grant Coordinators can import grants. "
                       "See the Nomad Compass User Manual for full details.")
ws[f"A{row}"].font = Font(size=10, italic=True, color=GREY)
ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")

# ------------------------------------------------------------------- Grants
ws2 = wb.create_sheet("Grants")
for i, w in enumerate(COL_WIDTHS, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

header_fill = PatternFill("solid", fgColor=TEAL)
header_font = Font(bold=True, color=WHITE, size=11)
for col, h in enumerate(HEADERS, start=1):
    c = ws2.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border_all
ws2.row_dimensions[1].height = 30
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = "A1:J51"
ws2.sheet_properties.pageSetUpPr.fitToPage = True
ws2.sheet_properties.pageSetUpPr.orientation = "landscape"

ex_font = Font(italic=True, color=GREY, size=11)
ex_fill = PatternFill("solid", fgColor=EXAMPLE_BG)
r = 2
for ex in EXAMPLES:
    for col, val in enumerate(ex, start=1):
        c = ws2.cell(row=r, column=col, value=val)
        c.font = ex_font
        c.fill = ex_fill
        c.border = border_all
        c.alignment = Alignment(horizontal="left" if col <= 2 else "center", vertical="center")
        if col in (3, 4):
            c.number_format = '#,##0'
    r += 1

data_font = Font(size=11, color=INK)
for row_idx in range(r, 52):
    fill = PatternFill("solid", fgColor=ROW_ALT if row_idx % 2 == 0 else WHITE)
    for col in range(1, 11):
        c = ws2.cell(row=row_idx, column=col)
        c.font = data_font
        c.fill = fill
        c.border = border_all
        c.alignment = Alignment(horizontal="left" if col <= 2 else "center", vertical="center")
        if col in (3, 4):
            c.number_format = '#,##0'

# Data validation: money >= 0; enums via dropdowns
dv_money = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                          allow_blank=True, showErrorMessage=True, errorTitle="Numbers only",
                          error="Amounts must be numbers (0 or more). No $ or commas needed.")
dv_money.add("C5:C51")
dv_money.add("D5:D51")
ws2.add_data_validation(dv_money)

dv_status = DataValidation(type="list", formula1='"active,pending,completed"', allow_blank=True,
                           showErrorMessage=True, errorTitle="Pick from the list",
                           error="Status must be active, pending, or completed.")
dv_status.add("G5:G51")
ws2.add_data_validation(dv_status)

dv_restriction = DataValidation(type="list", formula1='"restricted,unrestricted"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Pick from the list",
                                error="Restriction must be restricted or unrestricted.")
dv_restriction.add("H5:H51")
ws2.add_data_validation(dv_restriction)

dv_freq = DataValidation(type="list", formula1='"monthly,quarterly,semiannual,annual,none"', allow_blank=True,
                         showErrorMessage=True, errorTitle="Pick from the list",
                         error="Report Frequency must be monthly, quarterly, semiannual, annual, or none.")
dv_freq.add("I5:I51")
ws2.add_data_validation(dv_freq)

ws2.merge_cells("A53:J53")
ws2["A53"].value = ("EXAMPLE rows are placeholders — delete or overwrite them. Then File → Save As → "
                    "CSV UTF-8 (*.csv), and upload under Grant Tracking → Bulk Grant Upload.")
ws2["A53"].font = Font(size=10, italic=True, color=GREY)
ws2["A53"].alignment = Alignment(wrap_text=True, vertical="top")
ws2.row_dimensions[53].height = 36

os.makedirs(OUT_DIR, exist_ok=True)
wb.save(OUT_FILE)
print("saved:", OUT_FILE)
