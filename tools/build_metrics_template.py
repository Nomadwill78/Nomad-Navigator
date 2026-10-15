#!/usr/bin/env python3
"""Builds the Nomad Compass Program Metrics Excel upload template.

Output: ~/workspace/your_files/nomad-compass/nomad-compass-program-metrics-template.xlsx

The workbook has two sheets:
  1. Instructions - how to fill the template, save as CSV, and upload it.
  2. Program Metrics - the fillable table, matching the CSV importer in
     components/DataManagementView.tsx (fuzzy header mapping targets):
       Program Name, Month, Beneficiaries Served, Total Budget Spent, Cost Per Capita
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = os.path.expanduser("~/workspace/your_files/nomad-compass")
OUT_FILE = os.path.join(OUT_DIR, "nomad-compass-program-metrics-template.xlsx")

TEAL = "0F766E"
TEAL_DARK = "115E59"
TEAL_LIGHT = "CCFBF1"
BRASS = "B45309"
BRASS_LIGHT = "FEF3C7"
INK = "1F2937"
GREY = "6B7280"
WHITE = "FFFFFF"
ROW_ALT = "F9FAFB"
EXAMPLE_BG = "F3F4F6"

thin = Side(style="thin", color="D1D5DB")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADERS = [
    "Program Name",
    "Month",
    "Beneficiaries Served",
    "Total Budget Spent",
    "Cost Per Capita",
]
COL_WIDTHS = [34, 16, 22, 20, 18]

EXAMPLES = [
    ["Youth Workforce Training", "Jan 2026", 1500, 18000, 12.00],
    ["Sanitation Skills Workshop", "Feb 2026", 1200, 14000, 11.67],
    ["Hygiene Education Outreach", "Mar 2026", 2000, 19500, 9.75],
]

INSTRUCTIONS_BODY = [
    ("What this file is for",
     "This template feeds the Program Metrics section of Nomad Compass. Fill it in, save it as a "
     "CSV file, and upload it under Manage Data → CSV Data Stream Integration. Each row becomes one "
     "program-month record on your Impact Overview dashboard."),
    ("Step 1 — Fill in the Program Metrics sheet",
     "• Replace or delete the three grey EXAMPLE rows — they are placeholders only.\n"
     "• One row per program per month (e.g. a program running all year gets 12 rows).\n"
     "• Program Name: the name of your program, exactly as you want it on the dashboard. (Required)\n"
     "• Month: the reporting month, e.g. Jan 2026, 2026-01, or Q1 2026. (Optional but recommended)\n"
     "• Beneficiaries Served: how many people the program reached that month — a plain number. (Required)\n"
     "• Total Budget Spent: what the program cost that month — a plain number, no $ or commas. (Required)\n"
     "• Cost Per Capita: optional. Leave blank and Nomad Compass calculates it for you "
     "(Total Budget Spent ÷ Beneficiaries Served)."),
    ("Step 2 — Save as CSV",
     "In Excel: File → Save As → choose “CSV UTF-8 (Comma delimited) (*.csv)”.\n"
     "In Google Sheets: File → Download → Comma Separated Values (.csv).\n"
     "The file must end in .csv — that is the only format the uploader accepts."),
    ("Step 3 — Upload to Nomad Compass",
     "1. Sign in and open Manage Data from the sidebar.\n"
     "2. In the “CSV Data Stream Integration” panel, drag your .csv file in (or click to browse).\n"
     "3. Check the column mapping — Nomad Compass guesses it automatically; fix any dropdown that "
     "looks wrong. Name, People Served and Total Cost must be mapped.\n"
     "4. Choose Append to add to existing data, or Overwrite to replace it.\n"
     "5. Review the preview table, then click Confirm & Stream Data."),
    ("Tips for clean uploads",
     "• Numbers only in the number columns — no dollar signs, commas, or text like “approx”.\n"
     "• Keep Program Names consistent month to month (“Youth Workforce” vs “Youth Workforce Training” "
     "will show as two different programs).\n"
     "• One header row at the top, no merged cells, no extra title rows above the headers.\n"
     "• You can rename the column headers if you like — Nomad Compass matches them by meaning, "
     "but the template headers work as-is."),
]

wb = Workbook()

# ---------------------------------------------------------------- Instructions
ws = wb.active
ws.title = "Instructions"
ws.sheet_properties.pageSetUpPr.fitToPage = True

ws.merge_cells("A1:A1")
title_cell = ws["A1"]
title_cell.value = "Nomad Compass — Program Metrics Upload Template"
title_cell.font = Font(size=18, bold=True, color=WHITE)
title_cell.fill = PatternFill("solid", fgColor=TEAL)
title_cell.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 42
ws.column_dimensions["A"].width = 110

row = 3
for heading, body in INSTRUCTIONS_BODY:
    ws.merge_cells(f"A{row}:A{row}")
    h = ws[f"A{row}"]
    h.value = heading
    h.font = Font(size=13, bold=True, color=TEAL_DARK)
    h.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26
    row += 1
    ws.merge_cells(f"A{row}:A{row}")
    b = ws[f"A{row}"]
    b.value = body
    b.font = Font(size=11, color=INK)
    b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    # height heuristic: ~ one line per 100 chars at 110 width
    ws.row_dimensions[row].height = max(60, 18 * (len(body) // 100 + body.count("\n") + 1))
    row += 2

ws.merge_cells(f"A{row}:A{row}")
note = ws[f"A{row}"]
note.value = ("Questions? Your grant coordinator or impact analyst can edit metrics; other roles "
              "may be read-only. See the Nomad Compass User Manual for full details.")
note.font = Font(size=10, italic=True, color=GREY)
note.alignment = Alignment(wrap_text=True, vertical="top")
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------- Program Metrics
ws2 = wb.create_sheet("Program Metrics")
for i, w in enumerate(COL_WIDTHS, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

header_fill = PatternFill("solid", fgColor=TEAL)
header_font = Font(bold=True, color=WHITE, size=11)
for col, h in enumerate(HEADERS, start=1):
    c = ws2.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border_all
ws2.row_dimensions[1].height = 30
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = f"A1:E51"
ws2.sheet_properties.pageSetUpPr.fitToPage = True

# Example rows (greyed, clearly placeholders)
ex_font = Font(italic=True, color=GREY, size=11)
ex_fill = PatternFill("solid", fgColor=EXAMPLE_BG)
r = 2
for ex in EXAMPLES:
    for col, val in enumerate(ex, start=1):
        c = ws2.cell(row=r, column=col, value=val)
        c.font = ex_font
        c.fill = ex_fill
        c.border = border_all
        c.alignment = Alignment(horizontal="left" if col <= 2 else "right", vertical="center")
        if col >= 3:
            c.number_format = '#,##0' if col == 3 else '#,##0.00'
    r += 1

# Blank formatted rows for data entry
data_font = Font(size=11, color=INK)
for row_idx in range(r, 52):
    fill = PatternFill("solid", fgColor=ROW_ALT if row_idx % 2 == 0 else WHITE)
    for col in range(1, 6):
        c = ws2.cell(row=row_idx, column=col)
        c.font = data_font
        c.fill = fill
        c.border = border_all
        c.alignment = Alignment(horizontal="left" if col <= 2 else "right", vertical="center")
        if col == 3:
            c.number_format = '#,##0'
        elif col >= 4:
            c.number_format = '#,##0.00'

# Data validation: whole numbers for beneficiaries, decimals for money
dv_served = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                           allow_blank=True, showErrorMessage=True,
                           errorTitle="Numbers only",
                           error="Beneficiaries Served must be a whole number (0 or more).")
dv_served.add("C5:C51")
ws2.add_data_validation(dv_served)

dv_money = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Numbers only",
                          error="Money columns must be numbers (0 or more). No $ or commas needed.")
dv_money.add("D5:D51")
dv_money.add("E5:E51")
ws2.add_data_validation(dv_money)

# A note row under the table
ws2.merge_cells("A53:E53")
footer = ws2["A53"]
footer.value = ("EXAMPLE rows are placeholders — delete or overwrite them. Leave “Cost Per Capita” blank "
                "to have Nomad Compass calculate it. Then File → Save As → CSV UTF-8 (*.csv).")
footer.font = Font(size=10, italic=True, color=GREY)
footer.alignment = Alignment(wrap_text=True, vertical="top")
ws2.row_dimensions[53].height = 36

os.makedirs(OUT_DIR, exist_ok=True)
wb.save(OUT_FILE)
print("saved:", OUT_FILE)
